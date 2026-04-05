import json
import os
import io
import pypdf
import docx
from pptx import Presentation
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from openai import OpenAI
from dotenv import load_dotenv
from sqlalchemy.orm import Session
from fastapi import Depends
from database import get_db
from db import models
import schemas
import requests

# Load the .env file
load_dotenv()

# Initialize the router
router = APIRouter(tags=["Quiz Generation"])

# Get the key from the hidden file
api_key = os.getenv("OPENAI_API_KEY")
if not api_key:
    raise ValueError("No OpenAI API Key found. Check your .env file!")

# Initialize the client using the environment variable
client = OpenAI(api_key=api_key)

# --- YOUR SHARED HELPER FUNCTION ---
async def generate_questions_with_ai(extracted_text: str, config_data: list):
    prompt = f"""
    You are an educational assistant. Generate a quiz based ONLY on the text provided below.
    
    TEXT CONTENT:
    {extracted_text[:10000]} 

    CONFIGURATIONS:
    {json.dumps(config_data)} 

    OUTPUT FORMAT:
    Return a JSON object with a single key "questions". 
    Each question object MUST have:
    - "type": (e.g., "MCQ", "Essay", or "True/False")
    - "question": The text of the question.
    - "options": A list of strings (4 for MCQ, 2 for True/False, null for Essay).
    - "answer": The string text of the correct answer.
    - "explanation": A brief explanation.
    - "general_feedback": A detailed model answer or explanation to be shown to the student after submission.
    """

    try:
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "system", "content": "You are a quiz generator that outputs strictly valid JSON."},
                      {"role": "user", "content": prompt}],
            response_format={ "type": "json_object" }
        )
        
        quiz_data = json.loads(response.choices[0].message.content)
        
        if "questions" not in quiz_data:
            return []
            
        return quiz_data["questions"]
        
    except Exception as e:
        print(f"OpenAI Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/edit-question/")
async def edit_question(request: dict):
    old_q = request.get("question_data")
    instruction = request.get("instruction")

    prompt = f"""
    Modify this quiz question based on this instruction: {instruction}
    Original Question: {json.dumps(old_q)}
    
    Return ONLY the updated JSON object for the question.
    """

    try:
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            response_format={ "type": "json_object" }
        )
        return {"updated_question": json.loads(response.choices[0].message.content)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/generate-direct")
async def generate_direct(
    file: UploadFile = File(...),
    configs: str = Form(...)
):
    try:
        config_data = json.loads(configs)
        contents = await file.read()
        extracted_text = ""
        
        filename = file.filename.lower()
        
        if filename.endswith(".pdf"):
            pdf_reader = pypdf.PdfReader(io.BytesIO(contents))
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n"
                    
        elif filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(contents))
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
                
        elif filename.endswith(".pptx"):
            ppt = Presentation(io.BytesIO(contents))
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        extracted_text += shape.text + "\n"
                        
        elif filename.endswith(".txt"):
            extracted_text = contents.decode("utf-8")
            
        else:
            raise HTTPException(status_code=400, detail="Unsupported file. Please upload PDF, DOCX, PPTX, or TXT.")

        if not extracted_text.strip():
            raise HTTPException(status_code=400, detail="Could not extract text from the file.")

        questions = await generate_questions_with_ai(extracted_text, config_data)
        
        return {"questions": questions}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    
@router.get("/courses/", response_model=list[schemas.CourseResponse])
def get_courses_with_materials(db: Session = Depends(get_db)):
    courses = db.query(models.Course).all()
    return courses

@router.post("/courses/{course_id}/generate-quiz/")
async def generate_quiz(course_id: int, request: schemas.QuizGenerateRequest, db: Session = Depends(get_db)):
    selected_ids = request.selected_material_ids
    configs = request.configs
    
    # 1. Find the selected files in NeonDB
    materials = db.query(models.Material).filter(
        models.Material.course_id == course_id,
        models.Material.id.in_(selected_ids)
    ).all()
    
    if not materials:
        raise HTTPException(status_code=404, detail="No materials selected or found.")

    combined_text = ""
    for m in materials:
        try:
            # 2. Download the file from the Firebase URL directly into server RAM
            response = requests.get(m.file_path)
            response.raise_for_status() 
            file_bytes = io.BytesIO(response.content)
            
            filename = m.file_name.lower()
            
            # 3. Extract text from the downloaded bytes
            if filename.endswith(".pdf"):
                reader = pypdf.PdfReader(file_bytes)
                for page in reader.pages:
                    combined_text += page.extract_text() or ""
            elif filename.endswith(".docx"):
                doc = docx.Document(file_bytes)
                for para in doc.paragraphs:
                    combined_text += para.text + "\n"
            elif filename.endswith(".pptx"):
                ppt = Presentation(file_bytes)
                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            combined_text += shape.text + "\n"
            elif filename.endswith(".txt"):
                combined_text += file_bytes.read().decode("utf-8") + "\n"
                
        except Exception as e:
            print(f"Error reading cloud file {m.file_name}: {e}")

    if not combined_text.strip():
        raise HTTPException(status_code=400, detail="The selected files contain no readable text.")

    # 4. Send the extracted text to your AI helper!
    configs_as_dicts = [c.dict() for c in configs]
    questions = await generate_questions_with_ai(combined_text, configs_as_dicts)
    
    return {"questions": questions}