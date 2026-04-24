# backend/main.py

import os
import io
import json
import uuid
import shutil
import logging
import urllib.parse

import docx
import pypdf
import requests
from pptx import Presentation
from dotenv import load_dotenv

from fastapi import FastAPI, Depends, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session

# --- Internal Imports ---
from db import models
import schemas
from db.database import engine, get_db

from api.workspace_routes import router as workspace_router
from api.section_routes import router as section_router
from api.student_routes import router as student_router
from api.quiz_routes import router as quiz_router
from api.auth_routes import router as auth_router
from api.profile_routes import router as profile_router
from api.materials_routes import router as material_router

# --- External Services ---
from openai import OpenAI
import firebase_admin
from firebase_admin import credentials, storage


# ==============================================================================
# ── SETUP & INITIALIZATION ────────────────────────────────────────────────────
# ==============================================================================

# Logging Setup
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(name)s — %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

# Load Environment Variables
load_dotenv()

# OpenAI Setup
api_key = os.getenv("OPENAI_API_KEY")
if not api_key:
    raise ValueError("No OpenAI API Key found. Check your .env file!")
client = OpenAI(api_key=api_key)

# Firebase Setup
cred = credentials.Certificate("firebase_credentials.json")
firebase_admin.initialize_app(cred, {
    'storageBucket': os.getenv("FIREBASE_BUCKET")
})

# Ensure local upload directory exists
UPLOAD_DIR = "uploaded_materials"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# ==============================================================================
# ── FASTAPI APP CREATION ──────────────────────────────────────────────────────
# ==============================================================================

app = FastAPI(title="InstructorMate API", version="1.0.0")

# Combined CORS Middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def startup():
    print("✅ API started. (Database table creation is disabled).")

app.include_router(workspace_router)
app.include_router(section_router)
app.include_router(student_router)
app.include_router(quiz_router)
app.include_router(auth_router, prefix="/auth", tags=["Authentication"])
app.include_router(profile_router)
app.include_router(material_router)


# ==============================================================================
# ── HELPER FUNCTIONS ──────────────────────────────────────────────────────────
# ==============================================================================

async def generate_questions_with_ai(extracted_text: str, config_data: list):
    prompt = f"""
    You are an educational assistant. Generate a quiz based ONLY on the text provided below.
    
    TEXT CONTENT:
    {extracted_text[:10000]} 

    CONFIGURATIONS:
    {json.dumps(config_data)} 

    OUTPUT FORMAT:
    Return a JSON object with a single key "questions". 
    Each question object MUST have:
    - "type": (e.g., "MCQ", "Essay", or "True/False")
    - "question": The text of the question.
    - "options": A list of strings (4 for MCQ, 2 for True/False, null for Essay).
    - "answer": The string text of the correct answer.
    - "explanation": A brief explanation.
    - "general_feedback": A detailed model answer or explanation to be shown to the student after submission.
    """

    try:
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": "You are a quiz generator that outputs strictly valid JSON."},
                {"role": "user", "content": prompt}
            ],
            response_format={ "type": "json_object" }
        )
        
        quiz_data = json.loads(response.choices[0].message.content)
        
        if "questions" not in quiz_data:
            return []
            
        return quiz_data["questions"]
        
    except Exception as e:
        logger.error(f"OpenAI Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ==============================================================================
# ── ROUTES ────────────────────────────────────────────────────────────────────
# ==============================================================================

@app.get("/")
@app.get("/health")
def health():
    return {"status": "ok", "service": "InstructorMate API"}

# --- workspace & Material Routes ---

@app.post("/workspaces/", response_model=schemas.WorkspaceResponse)
def create_workspace(workspace: schemas.WorkspaceCreate, db: Session = Depends(get_db)):
    db_workspace = models.Workspace(**workspace.dict())
    db.add(db_workspace)
    db.commit()
    db.refresh(db_workspace)
    return db_workspace

@app.get("/workspaces/", response_model=list[schemas.WorkspaceResponse])
def get_workspaces_with_materials(db: Session = Depends(get_db)):
    return db.query(models.Workspace).all()

# --- AI Quiz Generation Routes ---

@app.post("/workspaces/{workspace_id}/generate-quiz/")
async def generate_quiz(workspace_id: int, request: schemas.QuizGenerateRequest, db: Session = Depends(get_db)):
    selected_ids = request.selected_material_ids
    configs = request.configs
    
    materials = db.query(models.Material).filter(
        models.Material.workspace_id == workspace_id,
        models.Material.material_id.in_(selected_ids)
    ).all()
    
    if not materials:
        raise HTTPException(status_code=404, detail="No materials selected or found.")

    combined_text = ""
    for m in materials:
        try:
            response = requests.get(m.file_path)
            response.raise_for_status() 
            file_bytes = io.BytesIO(response.content)
            filename = m.file_name.lower()
            
            if filename.endswith(".pdf"):
                reader = pypdf.PdfReader(file_bytes)
                for page in reader.pages:
                    combined_text += page.extract_text() or ""
            elif filename.endswith(".docx"):
                doc = docx.Document(file_bytes)
                for para in doc.paragraphs:
                    combined_text += para.text + "\n"
            elif filename.endswith(".pptx"):
                ppt = Presentation(file_bytes)
                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            combined_text += shape.text + "\n"
            elif filename.endswith(".txt"):
                combined_text += file_bytes.read().decode("utf-8") + "\n"
                
        except Exception as e:
            logger.error(f"Error reading cloud file {m.file_name}: {e}")

    if not combined_text.strip():
        raise HTTPException(status_code=400, detail="The selected files contain no readable text.")

    configs_as_dicts = [c.dict() for c in configs]
    
    questions = await generate_questions_with_ai(combined_text, configs_as_dicts)
    return {"questions": questions}


@app.post("/generate-direct")
async def generate_direct(
    file: UploadFile = File(...),
    configs: str = Form(...)
):
    try:
        config_data = json.loads(configs)
        contents = await file.read()
        extracted_text = ""
        filename = file.filename.lower()
        
        if filename.endswith(".pdf"):
            pdf_reader = pypdf.PdfReader(io.BytesIO(contents))
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n"
        elif filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(contents))
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
        elif filename.endswith(".pptx"):
            ppt = Presentation(io.BytesIO(contents))
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        extracted_text += shape.text + "\n"
        elif filename.endswith(".txt"):
            extracted_text = contents.decode("utf-8")
        else:
            raise HTTPException(status_code=400, detail="Unsupported file. Please upload PDF, DOCX, PPTX, or TXT.")

        if not extracted_text.strip():
            raise HTTPException(status_code=400, detail="Could not extract text from the file.")

        questions = await generate_questions_with_ai(extracted_text, config_data)
        return {"questions": questions}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/edit-question/")
async def edit_question(request: dict):
    old_q = request.get("question_data")
    instruction = request.get("instruction")

    prompt = f"""
    Modify this quiz question based on this instruction: {instruction}
    Original Question: {json.dumps(old_q)}
    
    Return ONLY the updated JSON object for the question.
    """
    try:
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            response_format={ "type": "json_object" }
        )
        return {"updated_question": json.loads(response.choices[0].message.content)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
